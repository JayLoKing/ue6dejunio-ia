"""Lectura unificada de hojas Excel, sin importar el formato.

Los registros 2026 son .xlsb (Excel binario) y los de años pasados .xlsx.
Esta capa abstrae esa diferencia y entrega siempre una grilla de celdas
(lista de listas) o un DataFrame, segun convenga.
"""

from __future__ import annotations

import logging
import warnings
from pathlib import Path

import pandas as pd

logger = logging.getLogger(__name__)

# openpyxl avisa de extensiones no soportadas (formato condicional, etc.) en
# estos Excel; es ruido inofensivo para nuestra lectura de valores.
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")


def _read_xlsb_sheet(path: Path, sheet: str) -> list[list]:
    from pyxlsb import open_workbook

    grid: list[list] = []
    with open_workbook(str(path)) as wb:
        if sheet not in wb.sheets:
            raise KeyError(f"Hoja '{sheet}' no existe en {path.name}. Hojas: {wb.sheets}")
        with wb.get_sheet(sheet) as ws:
            for row in ws.rows():
                grid.append([c.v for c in row])
    return grid


def _read_xlsx_sheet(path: Path, sheet: str) -> list[list]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise KeyError(f"Hoja '{sheet}' no existe en {path.name}. Hojas: {wb.sheetnames}")
    ws = wb[sheet]
    grid = [list(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    return grid


def read_sheet_grid(path: Path, sheet: str) -> list[list]:
    """Devuelve una hoja como grilla de celdas (lista de filas)."""
    suffix = path.suffix.lower()
    if suffix == ".xlsb":
        return _read_xlsb_sheet(path, sheet)
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx_sheet(path, sheet)
    raise ValueError(f"Formato no soportado: {path.suffix} ({path.name})")


def list_sheets(path: Path) -> list[str]:
    """Lista los nombres de hojas de un libro."""
    suffix = path.suffix.lower()
    if suffix == ".xlsb":
        from pyxlsb import open_workbook

        with open_workbook(str(path)) as wb:
            return list(wb.sheets)
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


def grid_to_df(grid: list[list]) -> pd.DataFrame:
    """Convierte una grilla cruda en DataFrame (sin encabezados)."""
    return pd.DataFrame(grid)


def cell(grid: list[list], row: int, col: int):
    """Acceso seguro a una celda (None si fuera de rango)."""
    if 0 <= row < len(grid) and 0 <= col < len(grid[row]):
        return grid[row][col]
    return None


def find_file(folder: Path, contains: str) -> Path | None:
    """Busca el primer archivo Excel cuyo nombre contiene `contains`
    (case-insensitive). Ignora archivos temporales de Excel (~$...)."""
    contains_low = contains.lower()
    for ext in ("*.xlsb", "*.xlsx", "*.xlsm"):
        for p in sorted(folder.glob(ext)):
            if p.name.startswith("~$"):
                continue
            if contains_low in p.name.lower():
                return p
    logger.warning("No se encontro archivo que contenga '%s' en %s", contains, folder.name)
    return None
